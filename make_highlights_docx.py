"""Gera Claude outputs/EJOR_highlights.docx (3-5 itens, <= 85 caracteres cada) a partir da secao 'Highlights' de submission_materials_draft.md.

    python make_highlights_docx.py
"""
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).parent
txt = (ROOT / "Claude outputs" / "submission_materials_draft.md").read_text(encoding="utf-8").split("## Highlights")[1].split("## Graphical")[0]
items = [re.sub(r"\s*\(\d+\)\s*$", "", m.group(1)).strip() for m in re.finditer(r"^\d+\.\s+(.*)$", txt, flags=re.M)]
assert 3 <= len(items) <= 5 and all(len(i) <= 85 for i in items), items
doc = Document()
doc.add_heading("Highlights", level=1)
for it in items:
    doc.add_paragraph(it, style="List Bullet")
out = ROOT / "Claude outputs" / "EJOR_highlights.docx"
doc.save(out)
print(len(items), "itens ->", out)
