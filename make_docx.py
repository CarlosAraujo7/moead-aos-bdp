"""Converte Claude outputs/manuscript_draft_v2.md em .docx (EJOR aceita Word em coluna unica; PDF nao serve como fonte).

Sem Pandoc: usa python-docx. Cobre o que o manuscrito usa: titulo, secoes (#, ##, ###), paragrafos com **negrito**, *italico* e `codigo`,
listas (-, 1.), tabelas em pipe, imagens ![legenda](caminho) e blocos de citacao (> ...), que sao OMITIDOS (caixa de status interna).
Uso:  python make_docx.py [entrada.md] [saida.docx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Claude outputs" / "manuscript_draft_v2.md"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else SRC.with_suffix(".docx")

TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_runs(par, text, size=None, base_bold=False):
    for part in TOKEN.split(text):
        if not part:
            continue
        bold, ital, code = base_bold, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, code = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, ital = part[1:-1], True
        r = par.add_run(part)
        r.bold, r.italic = bold, ital
        if code:
            r.font.name = "Consolas"
        if size:
            r.font.size = Pt(size)


def shade(cell, color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def add_table(doc, rows):
    header = [c.strip() for c in rows[0].strip().strip("|").split("|")]
    body = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows[2:]]
    ncol = len(header)
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8 if ncol >= 7 else 9
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.paragraphs[0].text = ""
        add_runs(c.paragraphs[0], h, size=size, base_bold=True)
        shade(c)
    for i, r in enumerate(body, start=1):
        for j in range(ncol):
            c = t.rows[i].cells[j]
            c.paragraphs[0].text = ""
            add_runs(c.paragraphs[0], r[j] if j < len(r) else "", size=size)
    doc.add_paragraph()


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    for lvl, sz in ((1, 16), (2, 13), (3, 11.5)):
        h = doc.styles[f"Heading {lvl}"]
        h.font.name = "Times New Roman"; h.font.size = Pt(sz); h.font.bold = True
    i, first_title = 0, True
    n_tab = n_img = 0
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
            continue
        if ln.startswith(">"):                        # caixa de status interna: omitida
            while i < len(lines) and lines[i].startswith(">"):
                i += 1
            continue
        if ln.startswith("# "):
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(p, ln[2:].strip(), size=16, base_bold=True)
                first_title = False
            else:
                doc.add_heading(ln[2:].strip(), level=1)
            i += 1
            continue
        m = re.match(r"^(#{2,4}) (.*)$", ln)
        if m:
            txt = re.sub(r"`\[[^`]*\]`", "", m.group(2)).strip()      # tira marcadores `[...]` dos titulos
            doc.add_heading(txt, level=min(len(m.group(1)) - 1, 3))
            i += 1
            continue
        if ln.startswith("*Working title.") or ln.startswith("*Working title"):
            i += 1
            continue
        if ln.lstrip().startswith("|"):
            blk = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                blk.append(lines[i]); i += 1
            if len(blk) >= 2:
                add_table(doc, blk); n_tab += 1
            continue
        mi = re.match(r"^!\[(.*)\]\((.*)\)\s*$", ln)
        if mi:
            cap, rel = mi.group(1), mi.group(2)
            img = (SRC.parent / rel)
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                n_img += 1
            p = doc.add_paragraph()
            add_runs(p, cap, size=9)
            i += 1
            continue
        if ln.startswith("```"):                         # bloco de codigo / algoritmo: monoespacado
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph()
                r = p.add_run(lines[i].rstrip())
                r.font.name = "Consolas"; r.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Inches(0.3)
                i += 1
            i += 1
            doc.add_paragraph()
            continue
        mb = re.match(r"^(\s*)([-*]|\d+\.) (.*)$", ln)
        if mb:
            if mb.group(2)[0].isdigit():                 # numeracao manual (evita continuar a numeracao de listas anteriores)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3); p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, mb.group(2) + " " + mb.group(3))
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, mb.group(3))
            i += 1
            continue
        p = doc.add_paragraph()
        add_runs(p, ln.strip())
        p.paragraph_format.space_after = Pt(6)
        i += 1
    doc.save(OUT)
    print(f"salvo: {OUT}  (tabelas: {n_tab}, figuras: {n_img}, paragrafos: {len(doc.paragraphs)})")


if __name__ == "__main__":
    main()
